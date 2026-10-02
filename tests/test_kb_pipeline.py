"""Knowledge base: SSRF-safe fetch, extraction, chunking, crawling."""
import io
import socket

import httpx
import pytest

from app.services.kb import crawl as crawl_mod
from app.services.kb.chunk import chunk_text, embedding_input
from app.services.kb.crawl import CrawlConfig, CrawlState, crawl_slice, in_scope, normalize_url
from app.services.kb.extract import ExtractError, extract_file, extract_html
from app.services.kb.fetch import FetchError, assert_public_url, fetch
from tests.kb_helpers import public_dns


# --- SSRF guard ------------------------------------------------------------------------

@pytest.mark.parametrize(
    "url",
    [
        "http://localhost/admin",
        "http://127.0.0.1/",
        "http://10.0.0.5/",
        "http://192.168.1.10/",
        "http://169.254.169.254/latest/meta-data/",  # cloud metadata
        "http://[::1]/",
        "ftp://example.com/file",
        "http://metadata.internal/",
    ],
)
def test_private_and_unsafe_urls_are_refused(url):
    with pytest.raises(FetchError):
        assert_public_url(url)


def test_public_url_allowed(monkeypatch):
    monkeypatch.setattr(socket, "getaddrinfo", public_dns)
    assert assert_public_url("https://example.com/page") == "https://example.com/page"


def test_redirect_to_private_address_is_blocked(monkeypatch):
    def dns(host, port, *a, **k):
        ip = "10.0.0.1" if host == "evil.test" else "93.184.216.34"
        return [(2, 1, 6, "", (ip, port or 80))]

    monkeypatch.setattr(socket, "getaddrinfo", dns)

    def handler(req):
        return httpx.Response(302, headers={"location": "http://evil.test/secret"})

    with httpx.Client(transport=httpx.MockTransport(handler)) as c, pytest.raises(FetchError):
        fetch("https://public.test/start", client=c)


def test_oversized_response_is_refused(monkeypatch):
    monkeypatch.setattr(socket, "getaddrinfo", public_dns)
    handler = lambda req: httpx.Response(200, content=b"x" * 5000, headers={"content-type": "text/html"})  # noqa: E731
    with httpx.Client(transport=httpx.MockTransport(handler)) as c, pytest.raises(FetchError):
        fetch("https://example.com/", client=c, max_bytes=1000)


# --- Extraction ------------------------------------------------------------------------

PAGE = """<html><head><title>AI Weekend Workshop | IT Talent Hub</title></head><body>
<nav><a href="/">Home</a> <a href="/courses/">Courses</a> <a href="/contact">Contact</a></nav>
<main><article>
<h1>AI Weekend Workshop</h1>
<p>The AI Weekend Workshop is a practical introductory programme for people who want to use modern AI tools at work.
It runs online over one weekend, Saturday and Sunday, from 10am to 4pm each day.</p>
<h2>Fees</h2>
<p>The fee is £99 per participant, including all course materials and a certificate of completion for every attendee.</p>
<p>See also <a href="/courses/data-engineering">Data Engineering</a>.</p>
</article></main>
<footer>Copyright 2026 IT Talent Hub. All rights reserved. Cookie settings.</footer>
</body></html>"""


def test_html_main_content_keeps_headings_and_drops_chrome():
    out = extract_html(PAGE, "https://ittalenthub.co.uk/workshop")
    assert "£99 per participant" in out.text
    assert "# AI Weekend Workshop" in out.text or "AI Weekend Workshop" in out.text
    assert "## Fees" in out.text
    assert "Cookie settings" not in out.text
    assert "AI Weekend Workshop" in out.title
    assert "https://ittalenthub.co.uk/courses/data-engineering" in out.links


def test_docx_extraction_keeps_heading_levels():
    import docx

    d = docx.Document()
    d.add_heading("Course Overview", level=1)
    d.add_paragraph("Master SQL, Power BI and Azure in two months.")
    buf = io.BytesIO()
    d.save(buf)
    out = extract_file(buf.getvalue(), "brochure.docx")
    assert "# Course Overview" in out.text and "Power BI" in out.text


def test_xlsx_rows_become_labelled_lines():
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Course", "Price"])
    ws.append(["Data Engineering", 1499])
    buf = io.BytesIO()
    wb.save(buf)
    out = extract_file(buf.getvalue(), "prices.xlsx")
    assert "Course: Data Engineering; Price: 1499" in out.text


def test_csv_and_text():
    assert "Course: AI; Fee: 99" in extract_file(b"Course,Fee\nAI,99\n", "fees.csv").text
    assert extract_file(b"hello knowledge", "notes.txt").text == "hello knowledge"


def test_unsupported_and_empty_files_rejected():
    with pytest.raises(ExtractError):
        extract_file(b"\x00\x01", "virus.exe")
    with pytest.raises(ExtractError):
        extract_file(b"   ", "empty.txt")


# --- Chunking --------------------------------------------------------------------------

def test_chunks_follow_headings_and_respect_size():
    text = "# Course\n\nIntro para.\n\n## Fees\n\n" + ("Fee details sentence. " * 200) + "\n\n## Schedule\n\nWeekends only."
    chunks = chunk_text(text, target=800, max_chars=1000, overlap=100)
    assert all(len(c.text) <= 1000 for c in chunks)
    fee_chunks = [c for c in chunks if c.heading == "Course > Fees"]
    assert len(fee_chunks) > 1  # long section split into several overlapping chunks
    # small sections survive — as their own chunk or folded (with their heading) into a neighbour
    everything = "\n".join(f"{c.heading}\n{c.text}" for c in chunks)
    assert "Intro para." in everything and "Weekends only." in everything
    assert "Schedule" in everything


def test_consecutive_chunks_overlap():
    paras = "\n\n".join(f"Paragraph {i} explains topic number {i} in detail." * 4 for i in range(12))
    chunks = chunk_text(paras, target=600, max_chars=900, overlap=150)
    assert len(chunks) >= 3
    tail_of_first = chunks[0].text[-60:]
    assert tail_of_first.split()[-1] in chunks[1].text


def test_embedding_input_has_context_header():
    c = chunk_text("# Fees\n\nThe fee is £99.")[0]
    assert embedding_input("Workshop page", c).startswith("Workshop page — Fees")


# --- Crawling --------------------------------------------------------------------------

def test_normalize_url_strips_tracking_and_fragments():
    assert normalize_url("HTTPS://Example.com:443/a//b?utm_source=x&b=2&a=1#top") == "https://example.com/a/b?a=1&b=2"


def test_scope_rules():
    cfg = CrawlConfig(start_url="https://www.site.test/courses/", exclude_paths=["/courses/old*"])
    assert in_scope("https://site.test/courses/ai", cfg)          # www ignored
    assert not in_scope("https://other.test/courses/ai", cfg)     # other site
    assert not in_scope("https://site.test/blog/x", cfg)          # outside include path
    assert not in_scope("https://site.test/courses/old-2019", cfg)  # excluded
    assert not in_scope("https://site.test/courses/logo.png", cfg)  # asset


def _site_transport():
    body = lambda title, extra="": (  # noqa: E731
        f"<html><head><title>{title}</title></head><body><main><h1>{title}</h1>"
        f"<p>{title} has plenty of useful detail about courses, fees, schedules and enrolment for learners. "
        f"It explains who each programme is for, what you will learn week by week, how the live sessions run, "
        f"and how to book a place. Recordings are available afterwards and support is offered by email.</p>"
        f"{extra}</main></body></html>"
    )
    pages = {
        "/robots.txt": ("text/plain", "User-agent: *\nDisallow: /private/\nSitemap: https://site.test/sitemap.xml"),
        "/sitemap.xml": ("application/xml", "<urlset><url><loc>https://site.test/a</loc></url></urlset>"),
        "/": ("text/html", body("Home", '<a href="/a">A</a><a href="/b">B</a><a href="/private/x">P</a><a href="https://elsewhere.test/">E</a>')),
        "/a": ("text/html", body("Page A", '<a href="/c">C</a>')),
        "/b": ("text/html", body("Page B")),
        "/c": ("text/html", body("Page C", '<a href="/d">D</a>')),
        "/d": ("text/html", body("Page D")),
        "/private/x": ("text/html", body("Secret")),
    }
    fetched = []

    def handler(req):
        fetched.append(req.url.path)
        if req.url.path in pages:
            ct, content = pages[req.url.path]
            return httpx.Response(200, text=content, headers={"content-type": ct})
        return httpx.Response(404, text="nope")

    return httpx.MockTransport(handler), fetched


def test_crawl_obeys_robots_sitemap_scope_and_depth(monkeypatch):
    monkeypatch.setattr(socket, "getaddrinfo", public_dns)
    transport, fetched = _site_transport()
    cfg = CrawlConfig(start_url="https://site.test/", max_pages=20, max_depth=2)
    state = CrawlState()
    with httpx.Client(transport=transport) as c:
        pages = list(crawl_slice(cfg, state, budget_seconds=60, client=c, sleep=lambda s: None))
    urls = {p.url for p in pages}
    assert {"https://site.test/", "https://site.test/a", "https://site.test/b", "https://site.test/c"} <= urls
    assert "https://site.test/private/x" not in urls and "/private/x" not in fetched  # robots.txt
    assert not any("elsewhere" in u for u in urls)  # other site
    assert "https://site.test/d" not in urls  # depth: / (0) → a (1, via sitemap) → c (2) → d (3) too deep
    assert state.done and state.stopped_reason == "complete"
    assert state.skipped_robots == 1


def test_crawl_respects_max_pages_and_resumes_across_slices(monkeypatch):
    monkeypatch.setattr(socket, "getaddrinfo", public_dns)
    transport, _ = _site_transport()
    cfg = CrawlConfig(start_url="https://site.test/", max_pages=3, max_depth=5)
    state = CrawlState()
    clock = {"t": 0.0}
    monkeypatch.setattr(crawl_mod.time, "monotonic", lambda: clock["t"])

    def tick(_s):
        clock["t"] += 10.0  # 10s between pages → a 5s slice ends after its 2nd page

    with httpx.Client(transport=transport) as c:
        first = list(crawl_slice(cfg, state, budget_seconds=5, client=c, sleep=tick))
        assert not state.done and state.frontier  # stopped for the slice, more to do
        restored = CrawlState.from_dict(state.to_dict())
        second = list(crawl_slice(cfg, restored, budget_seconds=600, client=c, sleep=tick))
    assert restored.done and restored.stopped_reason == "max_pages"
    assert restored.fetched == 3
    assert len({p.url for p in first + second}) == len(first + second)  # no page twice


@pytest.mark.parametrize(
    "ip,public",
    [
        ("185.158.133.1", True),
        ("64:ff9b::b99e:8501", True),     # NAT64 of 185.158.133.1 (DNS64 networks)
        ("64:ff9b::7f00:1", False),       # NAT64 wrapping 127.0.0.1
        ("64:ff9b::a00:1", False),        # NAT64 wrapping 10.0.0.1
        ("::ffff:127.0.0.1", False),      # IPv4-mapped loopback
        ("2002:7f00:1::1", False),        # 6to4 wrapping 127.0.0.1
        ("2606:4700:4700::1111", True),   # ordinary public IPv6
        ("fe80::1%en0", False),           # link-local with zone id
    ],
)
def test_ip_classification_unwraps_embedded_ipv4(ip, public):
    from app.services.kb.fetch import _ip_is_public

    assert _ip_is_public(ip) is public


def test_tiny_heading_sections_are_merged_not_scattered():
    page = "\n\n".join(f"## Card {i}\n\nShort line {i}." for i in range(8))
    chunks = chunk_text(page, target=1500, max_chars=2000, overlap=200)
    assert len(chunks) == 1  # eight scraps → one meaningful chunk
    assert "Card 0" in chunks[0].heading and "## Card" not in chunks[0].text.split("\n")[0]
    assert all(f"Short line {i}." in chunks[0].text for i in range(8))


def test_short_single_snippet_is_kept():
    chunks = chunk_text("# Fees\n\nThe fee is £99.")
    assert len(chunks) == 1 and chunks[0].text == "The fee is £99."
