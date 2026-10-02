"""Turn fetched pages and uploaded files into clean, heading-aware text.

HTML: main content only (navigation, footers and cookie banners dropped) as Markdown, so
headings survive for chunking; links are returned for the crawler. JavaScript-rendered sites
are out of scope — we read the HTML the server sends.
"""
from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from typing import Optional
from urllib.parse import urljoin, urldefrag

from app.config import settings

FILE_TYPES = {
    "pdf": "PDF",
    "docx": "Word",
    "txt": "Text",
    "md": "Markdown",
    "markdown": "Markdown",
    "csv": "CSV",
    "tsv": "TSV",
    "xlsx": "Excel",
    "html": "HTML",
    "htm": "HTML",
    "pptx": "PowerPoint",
}

MIN_PAGE_CHARS = 120  # thinner pages are boilerplate / empty shells


class ExtractError(Exception):
    """Content we can't read. Message is safe to show the tenant."""


@dataclass
class Extracted:
    title: str
    text: str
    links: list[str] = field(default_factory=list)


def _clean(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def file_ext(filename: str) -> str:
    name = (filename or "").lower().rsplit("/", 1)[-1]
    return name.rsplit(".", 1)[-1] if "." in name else ""


# --- HTML -------------------------------------------------------------------------------

def _html_links(html: str, base_url: str) -> list[str]:
    try:
        import lxml.html

        doc = lxml.html.fromstring(html)
    except Exception:
        return []
    out: list[str] = []
    seen: set[str] = set()
    for el in doc.xpath("//a[@href]"):
        href = (el.get("href") or "").strip()
        if not href or href.startswith(("mailto:", "tel:", "javascript:", "#")):
            continue
        absolute, _ = urldefrag(urljoin(base_url, href))
        if absolute not in seen:
            seen.add(absolute)
            out.append(absolute)
    return out


def _html_title(html: str) -> str:
    m = re.search(r"<title[^>]*>(.*?)</title>", html, re.S | re.I)
    if not m:
        return ""
    import html as htmllib

    return re.sub(r"\s+", " ", htmllib.unescape(m.group(1))).strip()[:200]


def extract_html(html: str, url: str = "") -> Extracted:
    import trafilatura

    text = trafilatura.extract(
        html,
        url=url or None,
        output_format="markdown",
        include_tables=True,
        include_links=False,
        include_images=False,
        favor_recall=True,
    ) or ""
    title = ""
    try:
        meta = trafilatura.extract_metadata(html, default_url=url or None)
        title = (getattr(meta, "title", None) or "").strip()
    except Exception:
        pass
    return Extracted(
        title=(title or _html_title(html) or url)[:200],
        text=_clean(text),
        links=_html_links(html, url) if url else [],
    )


# --- Files ------------------------------------------------------------------------------

def _pdf(data: bytes) -> str:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
        pages = [(p.extract_text() or "") for p in reader.pages]
    except Exception as exc:
        raise ExtractError("Couldn't read this PDF.") from exc
    return "\n\n".join(p.strip() for p in pages if p.strip())


def _docx(data: bytes) -> str:
    import docx

    try:
        d = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise ExtractError("Couldn't read this Word file.") from exc
    lines: list[str] = []
    for p in d.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        style = (p.style.name if p.style is not None else "").lower()
        if style.startswith("heading"):
            level = re.sub(r"\D", "", style) or "2"
            lines.append("#" * min(6, int(level)) + " " + t)
        else:
            lines.append(t)
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))
    return "\n\n".join(lines)


def _rows_to_text(rows: list[list[str]]) -> str:
    max_rows, max_cols = int(settings.KB_MAX_TABLE_ROWS), int(settings.KB_MAX_TABLE_COLS)
    rows = [r[:max_cols] for r in rows[: max_rows + 1]]
    if not rows:
        return ""
    header = [h.strip() for h in rows[0]]
    out: list[str] = []
    for r in rows[1:]:
        pairs = [f"{h or f'col{i + 1}'}: {v.strip()}" for i, (h, v) in enumerate(zip(header, r)) if str(v).strip()]
        if pairs:
            out.append("; ".join(pairs))
    return "\n".join(out) if out else "\n".join(" | ".join(r) for r in rows)


def _csv(data: bytes, delimiter: str) -> str:
    text = data.decode("utf-8-sig", errors="replace")
    return _rows_to_text([row for row in csv.reader(io.StringIO(text), delimiter=delimiter)])


def _xlsx(data: bytes) -> str:
    import openpyxl

    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise ExtractError("Couldn't read this Excel file.") from exc
    parts: list[str] = []
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            rows.append(["" if v is None else str(v) for v in row])
            if len(rows) > int(settings.KB_MAX_TABLE_ROWS):
                break
        body = _rows_to_text(rows)
        if body:
            parts.append(f"## {ws.title}\n{body}")
    return "\n\n".join(parts)


def _pptx(data: bytes) -> str:
    from pptx import Presentation

    try:
        prs = Presentation(io.BytesIO(data))
    except Exception as exc:
        raise ExtractError("Couldn't read this PowerPoint file.") from exc
    slides: list[str] = []
    for i, slide in enumerate(prs.slides, start=1):
        texts = [sh.text_frame.text.strip() for sh in slide.shapes if getattr(sh, "has_text_frame", False)]
        texts = [t for t in texts if t]
        if texts:
            slides.append(f"## Slide {i}\n" + "\n".join(texts))
    return "\n\n".join(slides)


def extract_file(data: bytes, filename: str, content_type: Optional[str] = None) -> Extracted:
    ext = file_ext(filename)
    ct = (content_type or "").lower()
    if not ext and "pdf" in ct:
        ext = "pdf"
    if ext not in FILE_TYPES:
        raise ExtractError(
            "Unsupported file type. Use " + ", ".join(sorted({f".{e}" for e in FILE_TYPES if e != "markdown"})) + "."
        )
    title = (filename or "upload").rsplit("/", 1)[-1][:200]
    if ext == "pdf":
        text = _pdf(data)
    elif ext == "docx":
        text = _docx(data)
    elif ext in ("csv", "tsv"):
        text = _csv(data, "\t" if ext == "tsv" else ",")
    elif ext == "xlsx":
        text = _xlsx(data)
    elif ext == "pptx":
        text = _pptx(data)
    elif ext in ("html", "htm"):
        page = extract_html(data.decode("utf-8", errors="replace"))
        if not page.text:
            raise ExtractError("No readable text found in this file.")
        return Extracted(title=page.title or title, text=page.text)
    else:
        text = data.decode("utf-8-sig", errors="replace")
    text = _clean(text)
    if not text:
        raise ExtractError("No readable text found in this file.")
    return Extracted(title=title, text=text)


def _decode(body: bytes, content_type: str) -> str:
    m = re.search(r"charset=([\w-]+)", content_type or "", re.I)
    try:
        return body.decode(m.group(1) if m else "utf-8", errors="replace")
    except LookupError:
        return body.decode("utf-8", errors="replace")


def extract_response(body: bytes, content_type: str, url: str) -> Extracted:
    """Extract a fetched URL (HTML page, PDF, or plain text)."""
    ct = (content_type or "").lower()
    if "pdf" in ct or url.lower().split("?")[0].endswith(".pdf"):
        text = _clean(_pdf(body))
        name = url.rstrip("/").rsplit("/", 1)[-1] or url
        return Extracted(title=name[:200], text=text)
    if ct.startswith("text/plain"):
        return Extracted(title=url[:200], text=_clean(_decode(body, ct)))
    if "html" in ct or not ct:
        return extract_html(_decode(body, ct), url)
    raise ExtractError(f"Unsupported content type: {ct.split(';')[0]}")
